"""Excel writer — produces paired ``.xltx`` + ``.xlsx`` per Winston-Lutz session.

Implements the ``wl-result-export`` spec:
- Per-session output folder at ``<output_root>/<MACHINE>/WL/<MACHINE>_WL_<RUNFOLDER>/``
- ``.xltx`` copied verbatim from ``templates/winston_lutz.xltx``
- ``.xlsx`` produced by loading the copy, populating 24 named cells +
  ``template_version``, and adding one per-image sheet per ``image_keys[i]``
- Sheet names sanitised to Excel constraints (≤31 chars, no forbidden chars)

Public API:
    - :func:`set_named_cell`
    - :func:`sanitise_sheet_name`
    - :func:`write_session_output`
"""

from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from core.result_types import WLAnalysisResult

logger = logging.getLogger(__name__)

#: Current template version stamp (should match the value in build_xltx_template.py).
TEMPLATE_VERSION = "2026-06"

#: Characters Excel forbids in sheet names.
_FORBIDDEN_SHEET_CHARS = re.compile(r"[\\\/\?\*\[\]:]")

#: Excel's maximum sheet name length.
_MAX_SHEET_NAME_LEN = 31


# ---------------------------------------------------------------------------
# Named-cell writer
# ---------------------------------------------------------------------------


def set_named_cell(wb: Workbook, name: str, value: Any) -> None:
    """Resolve a defined name to its target cell and set the value.

    Args:
        wb: An openpyxl ``Workbook`` with defined names.
        name: The defined name (e.g. ``"max_2d_cax_to_bb"``).
        value: The value to write.

    Raises:
        KeyError: If ``name`` is not a defined name in the workbook.
    """
    defined_names = wb.defined_names
    if name not in defined_names:
        raise KeyError(
            f"Defined name '{name}' not found in workbook. Ensure the xltx template defines it."
        )
    dn = defined_names[name]
    # ``destinations()`` yields (worksheet_title, cell_range) tuples
    for sheet_title, coord in dn.destinations:
        wb[sheet_title][coord] = value
        return  # only set the first destination
    raise KeyError(f"Defined name '{name}' has no destinations in the workbook.")


# ---------------------------------------------------------------------------
# Sheet name sanitisation
# ---------------------------------------------------------------------------


def sanitise_sheet_name(name: str) -> str:
    """Sanitise a string for use as an Excel sheet name.

    - Replace ``\\ / ? * [ ] :`` with ``_``
    - Truncate to 30 chars + ``…`` (total 31) if length exceeds 31

    Args:
        name: The raw sheet name candidate.

    Returns:
        A sheet-name-safe string of at most 31 characters.
    """
    cleaned = _FORBIDDEN_SHEET_CHARS.sub("_", name)
    if len(cleaned) <= _MAX_SHEET_NAME_LEN:
        return cleaned
    return cleaned[:30] + "…"


# ---------------------------------------------------------------------------
# Session output writer
# ---------------------------------------------------------------------------


def _serialize_for_excel(val: Any) -> Any:
    """Convert a value to a type openpyxl can write.

    - Primitives (str/int/float/bool/None) pass through
    - Dicts/lists/tuples are JSON-serialised to a string
    - Anything else is str()'d
    """
    import json

    if val is None or isinstance(val, (str, int, float, bool)):
        return val
    if isinstance(val, (dict, list, tuple)):
        try:
            return json.dumps(val, default=str)
        except (TypeError, ValueError):
            return str(val)
    return str(val)


def _session_folder(output_root: Path, machine_id: str, runfolder_name: str) -> Path:
    """Build the per-session output folder path: ``<root>/<MACHINE>/WL/<MACHINE>_WL_<RUNFOLDER>``."""
    return Path(output_root) / machine_id / "WL" / f"{machine_id}_WL_{runfolder_name}"


def write_session_output(
    result: WLAnalysisResult,
    output_root: Path,
    template_path: Path,
) -> Path:
    """Write paired ``.xltx`` + ``.xlsx`` for a Winston-Lutz session.

    Creates ``<output_root>/<MACHINE>/WL/<MACHINE>_WL_<RUNFOLDER>/`` if needed,
    copies the template as ``.xltx``, loads it, populates the 24 metric named
    cells + ``template_version``, adds one per-image sheet per
    ``image_keys[i]``, and saves the ``.xlsx``.

    Args:
        result: The analysis result to write.
        output_root: The output root (``config.output.root``).
        template_path: Path to ``templates/winston_lutz.xltx``.

    Returns:
        The path to the written ``.xlsx`` file.
    """
    runfolder_name = Path(result.runfolder_path).name
    session_dir = _session_folder(output_root, result.machine_id, runfolder_name)
    session_dir.mkdir(parents=True, exist_ok=True)

    base_name = f"{result.machine_id}_WL_{runfolder_name}"

    # 1. Copy the template as .xltx (verbatim)
    xltx_out = session_dir / f"{base_name}.xltx"
    shutil.copy(str(template_path), str(xltx_out))
    logger.info("Copied template to %s", xltx_out)

    # 2. Load the copy, populate named cells, add per-image sheets, save as .xlsx
    wb = load_workbook(str(xltx_out))

    # Populate the 24 metric named cells from result.summary
    from core.config import WL_METRIC_NAMES

    for canonical in WL_METRIC_NAMES:
        value = result.summary.get(canonical, 0.0)
        set_named_cell(wb, canonical, value)

    # Populate template_version
    set_named_cell(wb, "template_version", TEMPLATE_VERSION)

    # 3. Add per-image sheets (header-based table: field in A, value in B)
    for key, detail in zip(result.image_keys, result.image_details, strict=True):
        safe_name = sanitise_sheet_name(key)
        # Avoid duplicate sheet names (append index if collision)
        if safe_name in wb.sheetnames:
            idx = 1
            while f"{safe_name[:28]}_{idx}" in wb.sheetnames:
                idx += 1
            safe_name = f"{safe_name[:28]}_{idx}"
        ws = wb.create_sheet(title=safe_name)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 28
        ws["A1"] = "Field"
        ws["B1"] = "Value"
        for row_idx, (field, val) in enumerate(sorted(detail.items()), start=2):
            ws.cell(row=row_idx, column=1, value=str(field))
            # Serialize nested dicts/lists to string — openpyxl can't write them directly
            cell_val: Any = _serialize_for_excel(val)
            ws.cell(row=row_idx, column=2, value=cell_val)

    # 4. Save as .xlsx
    xlsx_out = session_dir / f"{base_name}.xlsx"
    wb.save(str(xlsx_out))
    logger.info("Wrote session output: %s", xlsx_out)

    # tolerance_mm is NOT written to the xlsx (per wl-result-export spec)
    return xlsx_out


__all__ = [
    "TEMPLATE_VERSION",
    "sanitise_sheet_name",
    "set_named_cell",
    "write_session_output",
]
