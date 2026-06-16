"""Excel writer — produces paired ``.xltx`` + ``.xlsx`` per CatPhan session.

Implements the ``cbct-result-export`` spec:
- Per-session output folder at
  ``<output_root>/<CATEGORY>/<DISPLAY>/<PYLINAC>/CatPhan/<MACHINE>_CP_<RUNFOLDER>/``
- ``.xltx`` copied verbatim from ``templates/catphan_504.xltx``
- ``.xlsx`` produced by loading the copy, populating 18 metric named cells +
  ``template_version``, and writing 4 per-CTP-module sheets

Public API:
    - :func:`write_cbct_session_output`
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font

from core.config import CATPHAN_SUMMARY_NAMES
from core.excel_helpers import set_named_cell
from core.result_types import CatPhanAnalysisResult
from core.session_io import build_session_folder, copy_template_to_session

logger = logging.getLogger(__name__)

#: Current CatPhan template version stamp (should match build_catphan_xltx_template.py).
TEMPLATE_VERSION = "2026-06-cp"


def write_cbct_session_output(
    result: CatPhanAnalysisResult,
    output_root: Path,
    template_path: Path,
) -> Path:
    """Write paired ``.xltx`` + ``.xlsx`` for a CatPhan session.

    Creates the session folder hierarchy if needed, copies the template as
    ``.xltx``, loads it, populates the 18 metric named cells +
    ``template_version``, writes 4 per-CTP-module sheets, and saves the
    ``.xlsx``.

    Args:
        result: The CatPhan analysis result to write.
        output_root: The machine's output root (``machine.output_root``).
        template_path: Path to ``templates/catphan_504.xltx``.

    Returns:
        The path to the written ``.xlsx`` file.
    """
    runfolder_name = Path(result.runfolder_path).name
    base_name = f"{result.machine_id}_CP_{runfolder_name}"

    # 1. Build session folder + copy template as .xltx
    session_dir = build_session_folder(
        output_root=output_root,
        machine_id=result.machine_id,
        module_dir="CatPhan",
        runfolder_name=runfolder_name,
        file_prefix="CP",
    )
    xltx_out = copy_template_to_session(template_path, session_dir, base_name)
    logger.info("Copied CatPhan template to %s", xltx_out)

    # 2. Load the copy, populate named cells + per-module sheets, save as .xlsx
    wb = load_workbook(str(xltx_out))

    # Populate the 18 metric named cells from result.summary
    for name in CATPHAN_SUMMARY_NAMES:
        value = result.summary.get(name, 0.0)
        set_named_cell(wb, name, value)

    # Populate template_version
    set_named_cell(wb, "template_version", TEMPLATE_VERSION)

    # 3. Write per-CTP-module sheets
    _write_ctp404_sheet(wb, result.ctp404)
    _write_ctp486_sheet(wb, result.ctp486)
    _write_ctp528_sheet(wb, result.ctp528)
    _write_ctp515_sheet(wb, result.ctp515)

    # 4. Save as .xlsx
    xlsx_out = session_dir / f"{base_name}.xlsx"
    wb.save(str(xlsx_out))
    logger.info("Wrote CatPhan session output: %s", xlsx_out)

    return xlsx_out


# ---------------------------------------------------------------------------
# Per-CTP-module sheet writers
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True)


def _write_header(ws: Any, headers: list[str], row: int = 1) -> None:
    """Write a header row with bold formatting."""
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = _HEADER_FONT


def _write_ctp404_sheet(wb: Any, ctp404: dict[str, Any]) -> None:
    """Write the CTP404 sheet: HU linearity table + geometry + thickness."""
    if "CTP404" not in wb.sheetnames:
        wb.create_sheet(title="CTP404")
    ws = wb["CTP404"]

    row = 1
    # HU linearity table (7 materials x name/nominal/measured/difference/stdev/passed)
    _write_header(
        ws, ["Material", "Nominal HU", "Measured HU", "Difference", "Stdev", "Passed"], row
    )
    row += 1
    hu_rois = ctp404.get("hu_rois", {})
    for _name, roi in sorted(hu_rois.items()):
        ws.cell(row=row, column=1, value=str(roi.get("name", _name)))
        ws.cell(row=row, column=2, value=float(roi.get("nominal_value", 0.0)))
        ws.cell(row=row, column=3, value=float(roi.get("value", 0.0)))
        ws.cell(row=row, column=4, value=float(roi.get("difference", 0.0)))
        ws.cell(row=row, column=5, value=float(roi.get("stdev", 0.0)))
        ws.cell(row=row, column=6, value=bool(roi.get("passed", False)))
        row += 1

    # Geometry summary
    row += 1
    ws.cell(row=row, column=1, value="Geometry").font = _HEADER_FONT
    row += 1
    ws.cell(row=row, column=1, value="avg_line_distance_mm")
    ws.cell(row=row, column=2, value=float(ctp404.get("avg_line_distance_mm", 0.0)))
    row += 1
    line_distances = ctp404.get("line_distances_mm", [])
    for i, dist in enumerate(line_distances):
        ws.cell(row=row, column=1, value=f"line_{i + 1}_distance_mm")
        ws.cell(row=row, column=2, value=float(dist))
        row += 1

    # Thickness info
    row += 1
    ws.cell(row=row, column=1, value="Thickness").font = _HEADER_FONT
    row += 1
    ws.cell(row=row, column=1, value="measured_slice_thickness_mm")
    ws.cell(row=row, column=2, value=float(ctp404.get("measured_slice_thickness_mm", 0.0)))
    row += 1
    ws.cell(row=row, column=1, value="thickness_num_slices_combined")
    ws.cell(row=row, column=2, value=int(ctp404.get("thickness_num_slices_combined", 0)))


def _write_ctp486_sheet(wb: Any, ctp486: dict[str, Any]) -> None:
    """Write the CTP486 sheet: uniformity ROIs + indices + NPS summary."""
    if "CTP486" not in wb.sheetnames:
        wb.create_sheet(title="CTP486")
    ws = wb["CTP486"]

    row = 1
    # Uniformity ROIs (5 regions)
    _write_header(ws, ["Region", "Value (HU)", "Nominal", "Difference", "Stdev", "Passed"], row)
    row += 1
    rois = ctp486.get("rois", {})
    for _name, roi in sorted(rois.items()):
        ws.cell(row=row, column=1, value=str(roi.get("name", _name)))
        ws.cell(row=row, column=2, value=float(roi.get("value", 0.0)))
        ws.cell(row=row, column=3, value=float(roi.get("nominal_value", 0.0)))
        ws.cell(row=row, column=4, value=float(roi.get("difference", 0.0)))
        ws.cell(row=row, column=5, value=float(roi.get("stdev", 0.0)))
        ws.cell(row=row, column=6, value=bool(roi.get("passed", False)))
        row += 1

    # Uniformity indices
    row += 1
    ws.cell(row=row, column=1, value="Uniformity Indices").font = _HEADER_FONT
    row += 1
    ws.cell(row=row, column=1, value="uniformity_index")
    ws.cell(row=row, column=2, value=float(ctp486.get("uniformity_index", 0.0)))
    row += 1
    ws.cell(row=row, column=1, value="integral_non_uniformity")
    ws.cell(row=row, column=2, value=float(ctp486.get("integral_non_uniformity", 0.0)))

    # NPS summary
    row += 1
    ws.cell(row=row, column=1, value="NPS Summary").font = _HEADER_FONT
    row += 1
    ws.cell(row=row, column=1, value="nps_avg_power")
    ws.cell(row=row, column=2, value=float(ctp486.get("nps_avg_power", 0.0)))
    row += 1
    ws.cell(row=row, column=1, value="nps_max_freq")
    ws.cell(row=row, column=2, value=float(ctp486.get("nps_max_freq", 0.0)))


def _write_ctp528_sheet(wb: Any, ctp528: dict[str, Any]) -> None:
    """Write the CTP528 sheet: MTF table (9 rows: 10%-90%)."""
    if "CTP528" not in wb.sheetnames:
        wb.create_sheet(title="CTP528")
    ws = wb["CTP528"]

    _write_header(ws, ["MTF (%)", "lp/mm"])
    mtf_lp_mm = ctp528.get("mtf_lp_mm", {})
    row = 2
    for pct in range(10, 91, 10):
        val = mtf_lp_mm.get(str(pct), 0.0)
        ws.cell(row=row, column=1, value=int(pct))
        ws.cell(row=row, column=2, value=float(val))
        row += 1


def _write_ctp515_sheet(wb: Any, ctp515: dict[str, Any]) -> None:
    """Write the CTP515 sheet: low contrast ROIs (variable rows)."""
    if "CTP515" not in wb.sheetnames:
        wb.create_sheet(title="CTP515")
    ws = wb["CTP515"]

    _write_header(
        ws,
        ["ROI", "Contrast", "CNR", "SNR", "Visibility", "Visibility Threshold", "Passed"],
    )
    roi_results = ctp515.get("roi_results", {})
    row = 2
    for _key, roi in sorted(roi_results.items()):
        ws.cell(row=row, column=1, value=str(_key))
        ws.cell(row=row, column=2, value=float(roi.get("contrast", 0.0)))
        ws.cell(row=row, column=3, value=float(roi.get("cnr", 0.0)))
        ws.cell(row=row, column=4, value=float(roi.get("signal to noise", 0.0)))
        ws.cell(row=row, column=5, value=float(roi.get("visibility", 0.0)))
        ws.cell(row=row, column=6, value=float(roi.get("visibility threshold", 0.0)))
        ws.cell(row=row, column=7, value=bool(roi.get("passed visibility", False)))
        row += 1


__all__ = [
    "TEMPLATE_VERSION",
    "write_cbct_session_output",
]
