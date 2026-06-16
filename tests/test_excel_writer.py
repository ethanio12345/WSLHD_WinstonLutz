"""Tests for ``core/excel_writer.py``."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.config import WL_METRIC_NAMES
from core.excel_writer import (
    sanitise_sheet_name,
    set_named_cell,
    write_session_output,
)
from core.result_types import WLAnalysisResult

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "winston_lutz.xltx"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_result(
    *,
    machine_id: str = "LA2",
    runfolder_path: str = "/data/LA2/WinstonLutz/2026-06-16_143022",
    n_images: int = 3,
) -> WLAnalysisResult:
    """Build a minimal WLAnalysisResult for writer testing."""
    image_keys = [f"G{i * 90}B0P0" for i in range(n_images)]
    image_details = [
        {
            "variable_axis": "Gantry" if i == 0 else "None",
            "gantry_angle": i * 90,
            "collimator_angle": 0,
            "couch_angle": 0,
            "cax2bb_distance": 0.5 + i * 0.1,
            "cax2epid_distance": 1.0 + i * 0.2,
        }
        for i in range(n_images)
    ]
    summary: dict[str, float | int | str] = {
        name: float(i) for i, name in enumerate(WL_METRIC_NAMES)
    }
    summary["machine_name"] = machine_id
    summary["session_date"] = "20260616"
    summary["num_total_images"] = n_images

    return WLAnalysisResult(
        machine_id=machine_id,
        runfolder_path=runfolder_path,
        session_date="20260616",
        summary=summary,
        image_details=image_details,
        image_keys=image_keys,
        params_used={"bb_size_mm": 5.0, "machine_scale": "VARIAN_IEC"},
    )


# ---------------------------------------------------------------------------
# set_named_cell
# ---------------------------------------------------------------------------


def test_set_named_cell_writes_value() -> None:
    """A defined name resolves to the correct cell."""
    wb = load_workbook(str(TEMPLATE_PATH))
    set_named_cell(wb, "max_2d_cax_to_bb", 1.23)

    # Resolve and verify
    dn = wb.defined_names["max_2d_cax_to_bb"]
    for sheet_title, coord in dn.destinations:
        assert wb[sheet_title][coord].value == 1.23
        return
    pytest.fail("max_2d_cax_to_bb has no destination")


def test_set_named_cell_missing_raises() -> None:
    """Missing defined name → KeyError with a clear message."""
    wb = load_workbook(str(TEMPLATE_PATH))
    with pytest.raises(KeyError, match="not found"):
        set_named_cell(wb, "bogus_name", 42)


# ---------------------------------------------------------------------------
# sanitise_sheet_name
# ---------------------------------------------------------------------------


def test_sanitise_safe_name_passthrough() -> None:
    """Safe names pass through unchanged."""
    assert sanitise_sheet_name("G0B0P0") == "G0B0P0"
    assert sanitise_sheet_name("G180B270P315") == "G180B270P315"


def test_sanitise_forbidden_chars() -> None:
    r"""Forbidden characters are replaced with ``_``."""
    assert sanitise_sheet_name("G0/B0:P0") == "G0_B0_P0"
    assert sanitise_sheet_name("test[1]") == "test_1_"
    assert sanitise_sheet_name("a\\b?c*d") == "a_b_c_d"


def test_sanitise_truncation() -> None:
    """Names >31 chars are truncated to 30 + ``…``."""
    long_name = "G180B270P315_extra_metadata_beyond_31_chars"
    result = sanitise_sheet_name(long_name)
    assert len(result) == 31
    assert result.endswith("…")
    assert result.startswith("G180B270P315_extra_metadata_be")


def test_sanitise_exactly_31_chars() -> None:
    """Names of exactly 31 chars are not truncated."""
    name_31 = "A" * 31
    assert sanitise_sheet_name(name_31) == name_31


# ---------------------------------------------------------------------------
# write_session_output
# ---------------------------------------------------------------------------


def test_write_session_output_creates_paired_files(tmp_path: Path) -> None:
    """Both .xltx and .xlsx are created in the session folder."""
    result = _make_result()
    xlsx_path = write_session_output(result, tmp_path, TEMPLATE_PATH)

    assert xlsx_path.exists()
    assert xlsx_path.suffix == ".xlsx"
    # The paired .xltx should also exist
    xltx_path = xlsx_path.with_suffix(".xltx")
    assert xltx_path.exists()


def test_write_session_output_folder_structure(tmp_path: Path) -> None:
    r"""Folder is ``<root>/<MACHINE>/WL/<MACHINE>_WL_<RUNFOLDER>``."""
    result = _make_result()
    xlsx_path = write_session_output(result, tmp_path, TEMPLATE_PATH)

    expected_dir = tmp_path / "LA2" / "WL" / "LA2_WL_2026-06-16_143022"
    assert xlsx_path.parent == expected_dir


def test_all_25_named_cells_populated(tmp_path: Path) -> None:
    """All 24 metric cells + template_version are populated."""
    result = _make_result()
    xlsx_path = write_session_output(result, tmp_path, TEMPLATE_PATH)
    wb = load_workbook(str(xlsx_path))

    from core.config import REQUIRED_TEMPLATE_NAMES

    for name in REQUIRED_TEMPLATE_NAMES:
        assert name in wb.defined_names, f"Missing defined name: {name}"
        dn = wb.defined_names[name]
        for sheet_title, coord in dn.destinations:
            value = wb[sheet_title][coord].value
            assert value is not None, f"Named cell {name} is empty"
            break


def test_sheet_count_is_n_plus_one(tmp_path: Path) -> None:
    """Sheet count = summary + N per-image sheets."""
    result = _make_result(n_images=5)
    xlsx_path = write_session_output(result, tmp_path, TEMPLATE_PATH)
    wb = load_workbook(str(xlsx_path))

    # 1 summary sheet + 5 per-image sheets
    assert len(wb.sheetnames) == 6


def test_per_image_sheet_titles_and_contents(tmp_path: Path) -> None:
    """Per-image sheets have the right title and contents (parallel alignment)."""
    result = _make_result(n_images=3)
    xlsx_path = write_session_output(result, tmp_path, TEMPLATE_PATH)
    wb = load_workbook(str(xlsx_path))

    per_image_sheets = wb.sheetnames[1:]  # skip summary
    assert per_image_sheets == ["G0B0P0", "G90B0P0", "G180B0P0"]

    # Check the first per-image sheet has header-based table
    ws = wb["G0B0P0"]
    assert ws["A1"].value == "Field"
    assert ws["B1"].value == "Value"
    # Should have rows for each field in image_details[0]
    detail = result.image_details[0]
    for row_idx in range(2, 2 + len(detail)):
        assert ws.cell(row=row_idx, column=1).value is not None


def test_tolerance_not_written(tmp_path: Path) -> None:
    """tolerance_mm is NOT written to any named cell or per-image field."""
    result = _make_result()
    xlsx_path = write_session_output(result, tmp_path, TEMPLATE_PATH)
    wb = load_workbook(str(xlsx_path))

    # No defined name should be 'tolerance_mm'
    assert "tolerance_mm" not in wb.defined_names

    # Check per-image sheets don't contain tolerance
    for sheet_name in wb.sheetnames[1:]:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            if row and isinstance(row[0], str) and "tolerance" in row[0].lower():
                pytest.fail(f"tolerance found in sheet {sheet_name}")


def test_rerun_overwrites_same_session(tmp_path: Path) -> None:
    """Re-running the same session overwrites both files."""
    result = _make_result()
    xlsx_path = write_session_output(result, tmp_path, TEMPLATE_PATH)
    first_mtime = xlsx_path.stat().st_mtime

    # Modify summary and re-write
    result2 = _make_result()
    object.__setattr__(result2, "summary", {**result.summary, "max_2d_cax_to_bb": 9.99})
    xlsx_path2 = write_session_output(result2, tmp_path, TEMPLATE_PATH)

    assert xlsx_path == xlsx_path2
    assert xlsx_path.stat().st_mtime >= first_mtime
