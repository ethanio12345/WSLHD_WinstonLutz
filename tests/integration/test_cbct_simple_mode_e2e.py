"""End-to-end test: CatPhan Simple mode pipeline (demo DICOMs -> xlsx).

Task 8.1 — Uses ``CatPhan504.from_demo_images()`` to populate a tmp_path
runfolder, writes a temporary ``machines.yaml`` with ``catphan`` configured,
runs ``run_cbct_analysis`` + ``write_cbct_session_output``, verifies the
produced xlsx exists at the expected deep path, has 19 named cells populated,
sheet count = 5.
"""

from __future__ import annotations

import warnings
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from core.cbct_excel_writer import write_cbct_session_output
from core.cbct_runner import run_cbct_analysis
from core.config import REQUIRED_CATPHAN_TEMPLATE_NAMES, load_config


@pytest.fixture(scope="module")
def catphan_runfolder(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Write the pylinac CatPhan504 demo DICOMs to a tmp runfolder."""
    runfolder = tmp_path_factory.mktemp("catphan_e2e") / "2026-06-16_monthly"
    runfolder.mkdir(parents=True)
    from pylinac import CatPhan504

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", FutureWarning)
        cbct = CatPhan504.from_demo_images()

    for i, img in enumerate(cbct.dicom_stack):
        img.save(str(runfolder / f"image_{i:04d}.dcm"))
    return runfolder


def test_cbct_simple_mode_e2e(catphan_runfolder: Path, tmp_path: Path) -> None:
    """Full Simple-mode pipeline: config -> analysis -> xlsx output."""
    # Set up the config pointing at the demo runfolder's parent as dicom_root
    dicom_root = catphan_runfolder.parent
    output_root = tmp_path / "output"
    output_root.mkdir()

    config_data = {
        "machines": {
            "LA2": {
                "display_name": "LA2 (TrueBeam)",
                "dicom_roots": {"catphan": str(dicom_root)},
                "output_root": str(output_root),
            }
        },
        "analysis_defaults": {
            "catphan": {
                "hu_tolerance": 40,
                "scaling_tolerance": 0.5,
                "slice_thickness_tolerance": 0.5,
            }
        },
        "assets": {"fry_meme_path": "/assets/fry.png", "logo_path": "/assets/logo.png"},
    }
    config_path = tmp_path / "machines.yaml"
    config_path.write_text(yaml.dump(config_data), encoding="utf-8")

    config = load_config(config_path)
    assert config.has_catphan()

    # Run the analysis
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        result = run_cbct_analysis(
            machine_id="LA2",
            runfolder_path=str(catphan_runfolder),
            hu_tolerance=40,
            scaling_tolerance=0.5,
            slice_thickness_tolerance=0.5,
            thickness_slice_straddle="auto",
            x_adjustment=0.0,
            y_adjustment=0.0,
            angle_adjustment=0.0,
            roi_size_factor=1.0,
            scaling_factor=1.0,
            minimum_rois_seen=3,
        )

    # Write the session output
    template_path = Path(__file__).resolve().parents[2] / "templates" / "catphan_504.xltx"
    xlsx_path = write_cbct_session_output(
        result=result,
        output_root=output_root,
        template_path=template_path,
    )

    # Verify the xlsx exists at the expected deep path
    assert xlsx_path.exists()
    path_str = str(xlsx_path)
    assert "CatPhan" in path_str
    assert "LA2_CP_2026-06-16_monthly" in path_str

    # Verify the paired xltx exists
    assert xlsx_path.with_suffix(".xltx").exists()

    # Verify the xlsx has all 19 named cells populated
    wb = load_workbook(str(xlsx_path))
    for name in REQUIRED_CATPHAN_TEMPLATE_NAMES:
        assert name in wb.defined_names, f"Defined name '{name}' missing"
        dn = wb.defined_names[name]
        for sheet_title, coord in dn.destinations:
            cell = wb[sheet_title][coord]
            assert cell.value is not None, f"Named cell '{name}' is empty"

    # Verify sheet count = 5
    assert len(wb.sheetnames) == 5
