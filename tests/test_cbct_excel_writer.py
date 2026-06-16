"""Tests for ``core/cbct_excel_writer.py`` — CatPhan Excel output.

Uses ``CatPhan504.from_demo_images()`` to produce a real analysis result,
then verifies the xlsx contract: 19 named cells populated, sheet count = 5,
per-module sheet contents correct, no ``tolerance_mm`` field anywhere.
"""

from __future__ import annotations

import warnings
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.cbct_excel_writer import TEMPLATE_VERSION, write_cbct_session_output
from core.config import REQUIRED_CATPHAN_TEMPLATE_NAMES
from core.result_types import CatPhanAnalysisResult

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def catphan_runfolder(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Write the pylinac CatPhan504 demo DICOMs to a tmp runfolder."""
    runfolder = tmp_path_factory.mktemp("catphan_demo")
    from pylinac import CatPhan504

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", FutureWarning)
        cbct = CatPhan504.from_demo_images()

    for i, img in enumerate(cbct.dicom_stack):
        img.save(str(runfolder / f"image_{i:04d}.dcm"))
    return runfolder


@pytest.fixture(scope="module")
def analysis_result(catphan_runfolder: Path) -> CatPhanAnalysisResult:
    """Run the analysis once and return the result."""
    from core.cbct_runner import run_cbct_analysis

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return run_cbct_analysis(
            machine_id="LA2",
            runfolder_path=str(catphan_runfolder),
            hu_tolerance=40,
            scaling_tolerance=0.5,
            slice_thickness_tolerance=0.5,
            thickness_slice_straddle="auto",
            x_adjustment=0.0,
            y_adjustment=0.0,
            angle_adjustment=0.0,
            roi_size_factor=1.0,
            scaling_factor=1.0,
            minimum_rois_seen=3,
        )


@pytest.fixture(scope="module")
def written_xlsx(
    analysis_result: CatPhanAnalysisResult,
    tmp_path_factory: pytest.TempPathFactory,
) -> Path:
    """Write the CatPhan session output and return the xlsx path."""
    output_root = tmp_path_factory.mktemp("cp_output")
    template_path = Path(__file__).resolve().parent.parent / "templates" / "catphan_504.xltx"
    return write_cbct_session_output(
        result=analysis_result,
        output_root=output_root,
        template_path=template_path,
    )


# ---------------------------------------------------------------------------
# Contract tests
# ---------------------------------------------------------------------------


def test_all_19_named_cells_populated(written_xlsx: Path) -> None:
    """Every one of the 19 defined names resolves to a non-empty cell."""
    wb = load_workbook(str(written_xlsx))
    for name in REQUIRED_CATPHAN_TEMPLATE_NAMES:
        assert name in wb.defined_names, f"Defined name '{name}' missing from workbook"
        dn = wb.defined_names[name]
        for sheet_title, coord in dn.destinations:
            cell = wb[sheet_title][coord]
            assert cell.value is not None, f"Named cell '{name}' at {sheet_title}!{coord} is empty"


def test_sheet_count_is_5(written_xlsx: Path) -> None:
    """The xlsx contains exactly 5 sheets: Summary + CTP404 + CTP486 + CTP528 + CTP515."""
    wb = load_workbook(str(written_xlsx))
    assert len(wb.sheetnames) == 5
    assert "Summary" in wb.sheetnames
    assert "CTP404" in wb.sheetnames
    assert "CTP486" in wb.sheetnames
    assert "CTP528" in wb.sheetnames
    assert "CTP515" in wb.sheetnames


def test_ctp404_has_7_materials(written_xlsx: Path) -> None:
    """CTP404 sheet contains the HU linearity table with all 7 materials."""
    wb = load_workbook(str(written_xlsx))
    ws = wb["CTP404"]
    # Header is row 1; data rows 2-8 (7 materials)
    materials = []
    for row in range(2, 9):
        val = ws.cell(row=row, column=1).value
        if val:
            materials.append(str(val))
    assert len(materials) == 7
    assert "Air" in materials
    assert "Teflon" in materials


def test_ctp486_has_5_regions(written_xlsx: Path) -> None:
    """CTP486 sheet contains the uniformity table with 5 regions."""
    wb = load_workbook(str(written_xlsx))
    ws = wb["CTP486"]
    # Header is row 1; data rows 2-6 (5 regions)
    regions = []
    for row in range(2, 7):
        val = ws.cell(row=row, column=1).value
        if val:
            regions.append(str(val))
    assert len(regions) == 5
    assert "Center" in regions


def test_ctp528_has_9_mtf_rows(written_xlsx: Path) -> None:
    """CTP528 sheet contains the MTF table with 9 rows (10%-90%)."""
    wb = load_workbook(str(written_xlsx))
    ws = wb["CTP528"]
    percentages = []
    for row in range(2, 11):
        val = ws.cell(row=row, column=1).value
        if val is not None:
            percentages.append(int(val))
    assert len(percentages) == 9
    assert percentages == [10, 20, 30, 40, 50, 60, 70, 80, 90]


def test_no_tolerance_mm_field(written_xlsx: Path) -> None:
    """No field named 'tolerance_mm' or similar appears in any sheet."""
    wb = load_workbook(str(written_xlsx))
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and isinstance(cell.value, str):
                    assert "tolerance_mm" not in cell.value.lower(), (
                        f"Found 'tolerance_mm' in {sheet_name}!{cell.coordinate}"
                    )


def test_template_version_stamped(written_xlsx: Path) -> None:
    """The template_version named cell contains the expected version."""
    wb = load_workbook(str(written_xlsx))
    dn = wb.defined_names["template_version"]
    for sheet_title, coord in dn.destinations:
        cell = wb[sheet_title][coord]
        assert cell.value == TEMPLATE_VERSION


def test_paired_xltx_exists(written_xlsx: Path) -> None:
    """The paired .xltx exists alongside the .xlsx."""
    xltx_path = written_xlsx.with_suffix(".xltx")
    assert xltx_path.exists()


def test_output_path_layout(written_xlsx: Path) -> None:
    """The output path follows the per-machine layout: <root>/CatPhan/LA2_CP_..."""
    path_str = str(written_xlsx)
    assert "CatPhan" in path_str
    assert "LA2_CP_" in written_xlsx.name
